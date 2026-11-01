"""Formato provisional sustituible cuando se aporten informes institucionales."""
from .comun import ROOT, leer

def texto_legible(text):
    import re
    labels={'representacion_acreditada':'acreditación de la representación',
            'fundamentacion_vinculada_al_acto':'fundamentación vinculada al acto',
            'informacion_financiera_completa':'información financiera de la rama',
            'solvencia_acreditada':'acreditación de solvencia',
            'identificacion_coincide':'concordancia de la identidad',
            'documentacion_registro_completa':'documentos completos de la rama de inscripción',
            'profesion_colegiacion_acreditada':'profesión y colegiación acreditadas',
            'condiciones_articulo25_revisadas':'condiciones del artículo 25 revisadas',
            'obligaciones_consorcio_identificadas':'obligaciones de cada integrante del consorcio',
            'datos_contratante':'datos completos del contratante',
            'rnp_vigente':'vigencia del RNP', 'informacion_actualizada':'información registral actualizada',
            'capacidad_legal_acreditada':'acreditación de la capacidad legal',
            'asistente_ia':'asistente de IA'}
    for key,value in sorted(labels.items(),key=lambda x:-len(x[0])):text=text.replace(key,value)
    text=re.sub(r'experiencia\.([A-Za-z0-9]+)\.',r'experiencia \1: ',text)
    return text.replace('_',' ')

def redactar(result, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    template=leer(ROOT/'templates/informe.v1.json')
    doc=Document()
    sec=doc.sections[0]
    sec.page_width,sec.page_height=Inches(8.5),Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.75)
    sec.left_margin=sec.right_margin=Inches(.8)
    for name,size in [('Normal',11),('Title',23),('Subtitle',11),('Heading 1',14),('Heading 2',12)]:
        style=doc.styles[name]
        style.font.name='Calibri'; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.core_properties.title=result['titulo']
    doc.core_properties.author='OECE Assistant'
    footer=sec.footer.paragraphs[0]
    footer.text='BORRADOR PARA REVISIÓN  |  '+result['id_ejecucion'][:12]+'  |  Página '
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
    for run in footer.runs: run.font.size=Pt(8)
    def table(headers, records, widths):
        t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
        for col,width in zip(t.columns,widths): col.width=Inches(width)
        for c,h,w in zip(t.rows[0].cells,headers,widths): c.text=h; c.width=Inches(w)
        header=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
        for record in records:
            for cell,value,width in zip(t.add_row().cells,record,widths):
                cell.text=str(value); cell.width=Inches(width)
        for rn,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ('top','left','bottom','right','insideH','insideV'):
                    b=OxmlElement('w:'+side)
                    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:'+k),v)
                    borders.append(b)
                pr.append(borders)
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7EEF5' if rn==0 else 'FFFFFF'); pr.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    m=OxmlElement('w:'+side); m.set(qn('w:w'),'100');m.set(qn('w:type'),'dxa');margins.append(m)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(4)
                    for r in p.runs: r.font.size=Pt(10);r.bold=rn==0
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    def url_link(p,url,label):
        element=OxmlElement('w:hyperlink')
        element.set(qn('r:id'),p.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
        run=OxmlElement('w:r');text=OxmlElement('w:t');text.text=label;run.append(text);element.append(run);p._p.append(element)
    doc.add_paragraph(result['titulo'],'Title')
    doc.add_paragraph(template['subtitulo'],'Subtitle')
    if result['sintetico']: doc.add_paragraph('EJEMPLO SINTÉTICO SIN DATOS DE UN PROVEEDOR REAL')
    doc.add_paragraph(result['resumen'])
    table(['Dato','Referencia'],[
        ('Expediente',result['expediente']),('Solicitante',result['solicitante']),
        ('Procedimiento',result['titulo']),('Fecha de solicitud declarada',result['fecha_solicitud'] or 'Por precisar'),
        ('Corte normativo',result['seleccion']['corte_normativo'])], [1.8,5.1])
    doc.add_heading(template['secciones']['objeto'],1)
    doc.add_paragraph('Examinar la documentación recibida, identificar las condiciones y evidencias aplicables y proponer observaciones para la evaluación del expediente. Este borrador no constituye una decisión del OECE ni acredita autenticidad documental o una consulta en línea a registros públicos.')
    doc.add_heading(template['secciones']['antecedentes'],1)
    for item in result['antecedentes']:
        doc.add_paragraph(item)
    doc.add_heading(template['secciones']['documentos'],1)
    table(['ID','Documento','Lectura'],[(d['id'],d['archivo'],str(len(d['fragmentos']))+' fragmentos'+(' con alertas' if any(b['alertas'] for b in d['fragmentos']) else '')) for d in result['documentos']], [1.25,4.2,1.45])
    doc.add_heading(template['secciones']['base'],1)
    regimes={'analisis_documental_por_procedimiento':'Evaluación documental del procedimiento seleccionado',
             'articulo_27_primera_ampliacion':'Artículo 27 para la primera ampliación posterior a la recategorización',
             'articulo_27_con_historial':'Artículo 27 con historial registral identificado','no_determinable':'Régimen pendiente de precisar'}
    doc.add_paragraph(regimes.get(result['seleccion']['regimen'],texto_legible(result['seleccion']['regimen']))+'. Los datos que definen la rama están declarados en la ficha del expediente. Las disposiciones seleccionadas se indican a continuación.')
    names={n['archivo']['sha256']:n['denominacion'] for n in leer(ROOT/'normativa/catalogo.json')['normas'] if n.get('archivo')}
    sources={}
    for row in result['matriz']:
        for src in [row.get('fuente_normativa'),*row.get('fuentes_complementarias',[])]:
            if not src: continue
            key=(src['archivo'],src['disposicion'],src['pagina_pdf'])
            if key not in sources: sources[key]=('N'+str(len(sources)+1),src)
    for source in result.get('normas_adicionales',[]):
        sources[(source['archivo'],source['disposicion'],source['pagina_pdf'])]=(source['codigo'],source)
    for code,src in sources.values():
        name=names.get(src['sha256'],src['archivo'])
        if src['archivo']=='tuo-27444-minjus-julio-2026.pdf':name='TUO de la Ley 27444, edición oficial MINJUSDH 2026'
        elif src['archivo']=='formula-cmc-spij.pdf':name='Folio oficial del Reglamento con la fórmula de CMC'
        p=doc.add_paragraph(f'{code}. {src["disposicion"]}; {name}, página PDF {src["pagina_pdf"]}. ')
        url_link(p,src['url'],'Fuente oficial')
    if not sources:
        doc.add_paragraph('Este módulo entrega un borrador documental asistido. Incorporar las disposiciones específicas verificadas en el análisis; no se reutilizan automáticamente las reglas de consultores.')
    doc.add_heading(template['secciones']['analisis'],1)
    counts=result['conteo']
    doc.add_paragraph(f'La matriz registra {counts.get("cumple",0)} verificaciones documentales satisfechas, {counts.get("no_cumple",0)} diferencias, {counts.get("no_determinable",0)} puntos no determinables y {counts.get("no_aplica",0)} puntos de otras ramas. Estos resultados se limitan a los hechos y documentos citados.')
    calculation=result.get('calculos',{}).get('cmc')
    if calculation:
        doc.add_heading('Cálculo propuesto de capacidad',2)
        if calculation['rama']=='con_experiencia':
            doc.add_paragraph(f'El monto neto atribuido de las obras es PEN {calculation["sumatoria_neta"]}. El coeficiente de gestión {calculation["coeficiente_gestion"]} resulta de sumar el factor de actividad {calculation["factor_actividad"]}, el financiero {calculation["factor_financiero"]} y el de ejecución {calculation["factor_ejecucion"]}. Al multiplicar el monto neto por ese coeficiente, la CMC propuesta es PEN {calculation["cmc_propuesta"]}.')
            doc.add_paragraph(f'El factor de ejecución utiliza la relación entre la suma neta atribuida y los montos contratados documentados, cuya suma es PEN {calculation["sumatoria_contratos"]}. Se deben revisar especialmente la base de cada importe y las obligaciones de consorcio antes de adoptar el cálculo.')
        else:
            doc.add_paragraph(f'La rama sin experiencia da una propuesta de PEN {calculation["cmc_propuesta"]}, según la información financiera disponible y la tabla aplicable. Esta cifra no constituye una capacidad otorgada ni acredita por sí sola el aumento solicitado.')
    calculation=result.get('calculos',{}).get('reconsideracion')
    if calculation:
        doc.add_heading('Cómputo propuesto del recurso',2)
        doc.add_paragraph(f'Con el calendario documentado se cuentan {calculation["dias_base"]} días hábiles y {calculation["dias_distancia_declarados"]} días de término de distancia declarado. El vencimiento calculado es {calculation["vencimiento"]}. El cómputo no acredita la eficacia de la notificación ni resuelve excepciones de horario, canal o régimen especial.')
    # Las explicaciones se mantienen en prosa; matriz compacta sin párrafos extensos.
    for row in result['matriz']:
        if row['resultado']=='no_aplica': continue
        doc.add_paragraph(row['id']+'  '+row['control'],'Heading 2')
        doc.add_paragraph(row['resultado'].replace('_',' ').capitalize()+'. '+texto_legible(row['motivo']))
        refs=row['evidencias']
        if refs:
            p=doc.add_paragraph('Sustento: '+'; '.join(dict.fromkeys(e['documento']+' '+e['ubicacion'] for e in refs))+'.')
            p.paragraph_format.keep_with_next=True
            quote=refs[0]['cita']
            p=doc.add_paragraph('Extracto: '+quote[:600]+('…' if len(quote)>600 else ''))
            p.paragraph_format.keep_with_next=bool(row['autores'])
        if row['autores']: doc.add_paragraph('Interpretación registrada por '+texto_legible(', '.join(row['autores']))+'.')
    for para in result.get('analisis',[]):
        doc.add_paragraph(para['texto'])
        doc.add_paragraph('Referencias: '+', '.join(para['referencias'])+'. Autor: '+para['autor']+'.')
    doc.add_heading(template['secciones']['observaciones'],1)
    if result['preguntas']:
        for q in result['preguntas']: doc.add_paragraph(texto_legible(q),'List Bullet')
    else: doc.add_paragraph('No se detectaron faltantes dentro de los controles ejecutados; verificar alcance y fuentes antes de adoptar una decisión administrativa.')
    doc.add_heading(template['secciones']['conclusion'],1)
    doc.add_paragraph(result['conclusion'])
    doc.add_heading(template['secciones']['anexos'],1)
    doc.add_paragraph('Acompañan este informe la matriz CSV y JSON, los originales conservados y el manifiesto de integridad. El formato es provisional y se ajustará a los modelos institucionales que aporte el usuario.')
    for d in result['documentos']:
        p=doc.add_paragraph(d['id']+' '+d['archivo']+'\nSHA256 '+d['sha256'])
        for r in p.runs: r.font.size=Pt(9)
    doc.save(path)
