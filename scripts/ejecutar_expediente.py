"""Circuito documental sintético. No decide cumplimiento ni ejecuta reglas legales."""
import argparse
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import shutil
import tempfile

from indexar_normativa import extract

ROOT = Path(__file__).resolve().parents[1]
VERSION = "0.1.0"


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8")


def load(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def inside(root, relative):
    path = (root / relative).resolve()
    if root not in path.parents:
        raise ValueError("Documento fuera de la carpeta del expediente")
    return path


def prepare(case_dir, rules_path):
    """Lee una instantánea comprobable y rechaza evidencia no localizable."""
    case_dir = case_dir.resolve()
    config_raw = (case_dir / "expediente.json").read_bytes()
    case = json.loads(config_raw)
    if case.get("sintetico") is not True:
        raise ValueError("Este prototipo admite únicamente expedientes sintéticos")
    rules_raw = rules_path.read_bytes()
    rules = json.loads(rules_raw)
    if rules.get("ejecutable") is not False or any(
        r["habilitada"] is not False or r["aprobacion_humana"] is not None
        for r in rules["reglas"]
    ):
        raise ValueError("El prototipo exige una matriz candidata deshabilitada")
    if case["procedimiento"] != rules["procedimiento_id"]:
        raise ValueError("Procedimiento incompatible")
    if not 1 <= len(case["documentos"]) <= 100:
        raise ValueError("Se requieren entre 1 y 100 documentos")
    documents, evidence, ids, hashes = [], {}, set(), {}
    for item in case["documentos"]:
        if item["id"] in ids:
            raise ValueError("Identificador documental duplicado")
        ids.add(item["id"])
        path = inside(case_dir, item["archivo"])
        raw = path.read_bytes()
        if len(raw) > 20_000_000:
            raise ValueError("Documento supera el límite de 20 MB")
        sha = digest(raw)
        if path.suffix.lower() == ".txt":
            engine, blocks, warnings = "utf8", [
                (f"linea:{n}", line, [])
                for n, line in enumerate(raw.decode("utf-8-sig").splitlines(), 1)
                if line.strip()
            ], []
        elif path.suffix.lower() in (".pdf", ".docx"):
            engine, blocks, warnings = extract(path)
        else:
            raise ValueError("Formato no admitido para expedientes: " + path.suffix)
        if digest(path.read_bytes()) != sha:
            raise ValueError("El documento cambió durante la lectura")
        def locator(loc):
            if isinstance(loc, str):
                return loc
            if "pagina_pdf" in loc:
                return "pagina_pdf:" + str(loc["pagina_pdf"])
            return loc["ruta_bloque"]
        fragments = [{"ubicacion": locator(loc), "texto": text, "alertas": flags}
                     for loc, text, flags in blocks]
        for fragment in fragments:
            evidence[(item["id"], fragment["ubicacion"])] = fragment
        documents.append({**item, "sha256": sha, "bytes": len(raw),
                          "lector": engine, "advertencias": warnings,
                          "duplicado_de": hashes.get(sha), "fragmentos": fragments})
        hashes.setdefault(sha, item["id"])
    known = {r["id"] for r in rules["reglas"]}
    mapping = case.get("vinculos", {})
    if set(mapping) - known:
        raise ValueError("Vínculo a control desconocido")
    rows = []
    for rule in rules["reglas"]:
        links = []
        for link in mapping.get(rule["id"], []):
            key = (link["documento"], link["ubicacion"])
            if key not in evidence or not evidence[key]["texto"].strip():
                raise ValueError("Evidencia inexistente o sin texto")
            source = next(d for d in documents if d["id"] == key[0])
            links.append({**link, "sha256": source["sha256"],
                          "texto": evidence[key]["texto"]})
        rows.append({"id": rule["id"], "rama": rule["rama"],
                     "control": rule["control"], "evidencias": links,
                     "resultado": "no_determinable", "fuente_normativa": rule["fuente"],
                     "motivo": "Control candidato pendiente de aprobación jurídica",
                     "limite": rule["limite"]})
    snapshot = {"version": VERSION, "implementacion_sha256": digest(Path(__file__).read_bytes()),
                "lector_sha256": digest((ROOT / "scripts/indexar_normativa.py").read_bytes()),
                "dependencias": {n: importlib.metadata.version(n) for n in ("pypdf", "python-docx")},
                "config_sha256": digest(config_raw),
                "matriz_sha256": digest(rules_raw),
                "documentos": [{k: d[k] for k in ("id", "archivo", "sha256")}
                               for d in documents]}
    run_id = digest(encoded(snapshot))
    return {"id_ejecucion": run_id, "sintetico": True,
            "expediente": case["id"], "solicitante": case["solicitante"],
            "procedimiento": case["procedimiento"], "fecha_solicitud": case["fecha_solicitud"],
            "estado": "borrador_pendiente_revision", "instantanea": snapshot,
            "documentos": documents, "matriz": rows,
            "conclusion": "No se determina cumplimiento ni categoría. Revisión humana pendiente."}


def report(result, path):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)
    for name, size in (("Normal", 11), ("Title", 23), ("Subtitle", 12), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    # Algunas plantillas de Word heredan una línea azul bajo Title.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    doc.core_properties.author = "OECE Assistant"
    doc.core_properties.title = "Revisión documental de ampliación de consultores"
    footer = sec.footer.paragraphs[0]
    footer.text = "DEMONSTRACIÓN SINTÉTICA  |  BORRADOR PARA REVISIÓN  |  "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.style = doc.styles["Normal"]
    for run in footer.runs:
        run.font.size = Pt(8)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
        for cell, value, width in zip(t.rows[0].cells, headers, widths):
            cell.text, cell.width = value, Cm(width)
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for c, value, width in zip(cells, row, widths):
                c.text, c.width = str(value), Cm(width)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for c in row.cells:
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = c._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E5E7EB" if i == 0 else "FFFFFF")
                pr.append(shade)
                borders = OxmlElement("w:tcBorders")
                for edge in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + edge)
                    for k, v in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        el.set(qn("w:" + k), v)
                    borders.append(el)
                pr.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    el = OxmlElement("w:" + edge)
                    el.set(qn("w:w"), "100")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                pr.append(margins)
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    for r in p.runs:
                        r.font.size = Pt(9)
                        r.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    sources = {}
    for row in result["matriz"]:
        source = row["fuente_normativa"]
        sources.setdefault(source["archivo"], (f"N{len(sources) + 1}", source))
    doc.add_paragraph("Revisión documental de ampliación de consultores", "Title")
    doc.add_paragraph("Demostración sintética para revisar el circuito de expediente", "Subtitle")
    doc.add_paragraph(f"Expediente {result['expediente']} · Solicitante {result['solicitante']}")
    doc.add_paragraph("No se determina cumplimiento ni categoría. Los 22 controles jurídicos siguen pendientes de aprobación. Este documento permite revisar la lectura, los vínculos de evidencia y la matriz antes de habilitar una evaluación legal.")
    doc.add_heading("Documentos recibidos", 1)
    table(["ID", "Archivo", "Lectura"],
          [(d["id"], d["archivo"], f"{len(d['fragmentos'])} fragmentos" +
            (f"; duplicado de {d['duplicado_de']}" if d["duplicado_de"] else ""))
           for d in result["documentos"]], [1.5, 10.8, 4.7])
    doc.add_heading("Alcance y revisión pendiente", 1)
    doc.add_paragraph("Los vínculos entre documentos y controles fueron declarados en la configuración de esta demostración; no los infirió un modelo. Una evidencia vinculada no acredita autenticidad, suficiencia ni cumplimiento. Las referencias normativas son candidatas y no certifican vigencia aplicable.")
    doc.add_paragraph("Revisar primero el régimen temporal y la representación. Después cotejar la experiencia acreditada con contratos, conformidades y pagos. Los controles sin evidencia quedan identificados en la matriz; su ausencia en este ejemplo no equivale a incumplimiento.")
    doc.add_paragraph(f"Fecha de solicitud declarada {result['fecha_solicitud']}. Ejecución {result['id_ejecucion'][:16]}.")
    doc.add_page_break()
    doc.add_heading("Matriz de revisión documental", 1)
    doc.add_paragraph("Todos los resultados son no determinables porque la matriz jurídica aún no está aprobada. N identifica la fuente del anexo; p indica la página del PDF.")
    table(["Control", "Punto de revisión", "Evidencia", "Fuente"],
          [(r["id"], r["control"], "; ".join(
              f"{e['documento']} {e['ubicacion']}" for e in r["evidencias"]
          ) or "Sin vínculo", sources[r["fuente_normativa"]["archivo"]][0] +
             f" p {r['fuente_normativa']['pagina_pdf']}")
           for r in result["matriz"]], [1.7, 9.4, 3.9, 2])
    doc.add_page_break()
    doc.add_heading("Trazabilidad de evidencias y fuentes", 1)
    doc.add_paragraph("Las huellas completas permiten comprobar qué versión fue leída. El archivo matriz.json conserva textos, localizadores, motivos y referencias por control.")
    for d in result["documentos"]:
        doc.add_paragraph(f"{d['id']}  {d['archivo']}", "Heading 2")
        p = doc.add_paragraph("SHA256 " + d["sha256"])
        p.runs[0].font.size = Pt(9)
        linked = [e for row in result["matriz"] for e in row["evidencias"] if e["documento"] == d["id"]]
        if linked:
            doc.add_paragraph(f"Extracto vinculado en {linked[0]['ubicacion']}: {linked[0]['texto']}")
    doc.add_heading("Fuentes normativas candidatas", 1)
    for name, (code, source) in sources.items():
        doc.add_paragraph(f"{code}  {name}", "Heading 2")
        p = doc.add_paragraph("SHA256 " + source["sha256"] + "  ")
        p.runs[0].font.size = Pt(9)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), p.part.relate_to(source["url"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Abrir original oficial"
        run.append(text)
        link.append(run)
        p._p.append(link)
    doc.save(path)


def execute(case_dir, output_dir, rules_path=ROOT / "rules/consultores/matriz.v0.1.json"):
    result = prepare(Path(case_dir), Path(rules_path))
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    final = output_dir / result["id_ejecucion"]
    if final.exists():
        manifest = load(final / "entrega.json")
        for name, sha in manifest["archivos"].items():
            if digest((final / name).read_bytes()) != sha:
                raise ValueError("Entrega previa alterada")
        return final, True
    temp = Path(tempfile.mkdtemp(prefix=".preparando-", dir=output_dir))
    try:
        (temp / "originales").mkdir()
        for n, item in enumerate(result["documentos"], 1):
            raw = inside(Path(case_dir).resolve(), item["archivo"]).read_bytes()
            if digest(raw) != item["sha256"]:
                raise ValueError("Original cambió antes de conservar la instantánea")
            (temp / "originales" / (str(n).zfill(3) + Path(item["archivo"]).suffix.lower())).write_bytes(raw)
        for source, name, expected in (
            (Path(case_dir) / "expediente.json", "expediente.json", result["instantanea"]["config_sha256"]),
            (Path(rules_path), "reglas.json", result["instantanea"]["matriz_sha256"]),
        ):
            raw = source.read_bytes()
            if digest(raw) != expected:
                raise ValueError("Configuración o matriz cambió antes de conservar la instantánea")
            (temp / name).write_bytes(raw)
        (temp / "matriz.json").write_bytes(encoded(result))
        report(result, temp / "informe.docx")
        manifest = {"id_ejecucion": result["id_ejecucion"],
                    "estados": ["recibido", "extraido", "borrador_pendiente_revision"],
                    "archivos": {p.relative_to(temp).as_posix(): digest(p.read_bytes())
                                 for p in sorted(temp.rglob("*")) if p.is_file()}}
        (temp / "entrega.json").write_bytes(encoded(manifest))
        os.rename(temp, final)
    finally:
        if temp.exists():
            if temp.resolve().parent != output_dir:
                raise ValueError("Carpeta temporal fuera del destino de salida")
            shutil.rmtree(temp)
    return final, False


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--entrada", type=Path, required=True)
    parser.add_argument("--salida", type=Path, required=True)
    args = parser.parse_args()
    folder, reused = execute(args.entrada, args.salida)
    print(json.dumps({"carpeta": str(folder), "reutilizada": reused}, ensure_ascii=True))
