import base64
from datetime import datetime,timezone
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.request import Request,urlopen
from urllib.error import HTTPError

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from oece.comun import guardar,leer,sha
from oece.evaluacion import evaluar,seleccionar,validar_hechos
from oece.flujo import Memoria,bloqueo,iniciar,preparar,procesar,verificar_entrega
from oece.coordinador import observar_una_vez,responder
from oece.lectura import leer_documento,aplicar_cotejos
from oece.panel import crear_servidor

def ficha():
    return {'id':'DEMO-V1','solicitante':'Proveedor Sintético','sintetico':True,
            'procedimiento':'ampliacion_consultores','fecha_solicitud':'2026-10-30','tipo_persona':'juridica',
            'origen_proveedor':'nacional','domicilio':'domiciliado','primera_ampliacion_tras_recategorizacion':True,
            'historial_rnp_disponible':True,'tramite_en_curso_transicion':False,'categoria_solicitada':'b',
            'experiencias':[{'id':'E1','origen':'peru_ley','consorcio':False}]}

def evidence_fixture(values):
    # Texto sintético con cada valor explícito, sin asumir que el nombre del campo
    # pruebe el dato en un expediente real. El autor es siempre IA de prueba.
    text='\n'.join(k+' = '+json.dumps(v,ensure_ascii=False) for k,v in values.items())
    digest=sha(text.encode())
    documents=[{'id':'D1','sha256':digest,'fragmentos':[{'ubicacion':'linea:1','texto':text,'alertas':[]}]}]
    facts=[{'campo':k,'valor':v,'autor':'Simulación de prueba','tipo_autor':'asistente_ia',
            'fecha':'2026-10-31T10:00:00-05:00','evidencias':[{'documento':'D1','sha256':digest,
            'ubicacion':'linea:1','cita':k+' = '+json.dumps(v,ensure_ascii=False)}]} for k,v in values.items()]
    return documents,facts

class CasosJuridicosV1(unittest.TestCase):
    def run_case(self,config,values):
        docs,facts=evidence_fixture(values)
        return {r['id']:r for r in evaluar(config,docs,facts)['matriz']}
    def test_S01_frontera_enero_no_se_resuelve_automaticamente(self):
        c=ficha();c['fecha_solicitud']='2026-01-19'
        rows=self.run_case(c,{'solicitud_presentada':True})
        self.assertEqual('no_determinable',rows['AC-01']['resultado'])
        self.assertEqual('no_determinable',rows['AC-03']['resultado'])
    def test_S02_categoria_antigua_sin_historial(self):
        c=ficha();c['historial_rnp_disponible']=False;c['categoria_anterior']='D'
        self.assertEqual('no_determinable',seleccionar(c)['regimen'])
    def test_S03_documento_equivalente_sin_titulo_constancia(self):
        rows=self.run_case(ficha(),{'experiencia.E1.documento_entidad':True,'experiencia.E1.datos_constancia_completos':True})
        self.assertEqual('cumple',rows['AC-06/E1']['resultado'])
    def test_S04_privada_con_sustento_faltante(self):
        c=ficha();c['experiencias'][0]['origen']='peru_fuera_ley'
        rows=self.run_case(c,{'experiencia.E1.contrato':True})
        self.assertEqual('no_determinable',rows['AC-07/E1']['resultado'])
        self.assertIn('datos_contratante',rows['AC-07/E1']['motivo'])
    def test_S05_diferencia_comprobantes_20000(self):
        c=ficha();c['experiencias'][0]['origen']='peru_fuera_ley'
        rows=self.run_case(c,{'experiencia.E1.monto_final':'100000.00','experiencia.E1.comprobantes':['80000.00'],
                              'experiencia.E1.moneda_monto':'PEN','experiencia.E1.moneda_comprobantes':'PEN'})
        self.assertEqual('no_cumple',rows['AC-08/E1']['resultado']);self.assertIn('20000',rows['AC-08/E1']['motivo'])
    def test_S06_consorcio_sin_obligaciones(self):
        c=ficha();c['experiencias'][0]['consorcio']=True
        self.assertEqual('no_determinable',self.run_case(c,{})['AC-10/E1']['resultado'])
    def test_S07_moneda_extranjera_sin_tasa(self):
        rows=self.run_case(ficha(),{'experiencia.E1.moneda_monto':'USD','experiencia.E1.monto_final':'10000'})
        self.assertEqual('no_determinable',rows['AC-15/E1']['resultado'])
    def test_S08_letra_historica_no_se_homologa(self):
        c=ficha();c['categoria_solicitada']='A'
        self.assertTrue(any('significado' in x for x in seleccionar(c)['preguntas']))
        self.assertEqual('no_determinable',self.run_case(c,{})['AC-11/E1']['resultado'])
    def test_S09_extranjero_sin_conformidad(self):
        c=ficha();c['experiencias'][0]['origen']='extranjero'
        self.assertEqual('no_determinable',self.run_case(c,{'experiencia.E1.contrato':True})['AC-09/E1']['resultado'])
    def test_S10_sin_revision_no_duplicar_experiencia(self):
        c=ficha();c['experiencias'].append({'id':'E2','origen':'peru_ley','consorcio':False})
        rows=self.run_case(c,{'experiencia.E1.objeto_principal_revisado':True})
        self.assertEqual('no_determinable',rows['AC-13/E1']['resultado'])
        self.assertEqual('no_determinable',rows['AC-13/E2']['resultado'])
    def test_extranjero_no_domiciliado_ruc_no_aplica(self):
        c=ficha();c.update(origen_proveedor='extranjero',domicilio='no_domiciliado')
        self.assertEqual('no_aplica',self.run_case(c,{})['AC-17']['resultado'])
    def test_fuente_cambiada_y_cita_inventada_se_rechazan(self):
        docs,facts=evidence_fixture({'solicitud_presentada':True});facts[0]['evidencias'][0]['cita']='Texto que no existe'
        with self.assertRaises(ValueError): validar_hechos(facts,docs)
        facts[0]['evidencias'][0]['cita']='solicitud_presentada';facts[0]['evidencias'][0]['sha256']='0'*64
        with self.assertRaises(ValueError): validar_hechos(facts,docs)
    def test_ocr_no_resuelve_hechos_favorables(self):
        docs,facts=evidence_fixture({'solicitud_presentada':True});docs[0]['fragmentos'][0]['alertas']=['ocr_sin_cotejo_visual']
        row=next(r for r in evaluar(ficha(),docs,facts)['matriz'] if r['id']=='AC-03')
        self.assertEqual('no_determinable',row['resultado'])
    def test_moneda_distinta_no_calcula_y_nan_no_pasa(self):
        c=ficha();c['experiencias'][0]['origen']='peru_fuera_ley'
        for amount in ('100','NaN','Infinity','-2'):
            rows=self.run_case(c,{'experiencia.E1.monto_final':amount,'experiencia.E1.comprobantes':['100'],
                                  'experiencia.E1.moneda_monto':'USD','experiencia.E1.moneda_comprobantes':'PEN'})
            self.assertEqual('no_determinable',rows['AC-08/E1']['resultado'])
    def test_condicion_ley32732_no_bloquea_art27(self):
        c=ficha()
        self.assertEqual('articulo_27_primera_ampliacion',seleccionar(c)['regimen'])
    def test_otros_modulos_no_declaran_evaluador_automatizado(self):
        for module in ('aumento_cmc','inscripcion','actualizacion','reconsideracion'):
            c=ficha();c['procedimiento']=module
            rows=self.run_case(c,{})
            self.assertEqual({'no_determinable'},{r['resultado'] for r in rows.values()})
    def test_inicio_anterior_no_rechaza_toda_experiencia(self):
        c=ficha();c['tipo_persona']='natural'
        rows=self.run_case(c,{'experiencia.E1.fecha_inicio_prestacion':'2020-01-01','experiencia.E1.fecha_colegiacion':'2020-02-01'})
        self.assertEqual('no_determinable',rows['AC-12/E1']['resultado'])
    def test_cotejo_ocr_conserva_original_y_rechaza_otra_imagen(self):
        docs,facts=evidence_fixture({'solicitud_presentada':True});b=docs[0]['fragmentos'][0];b.update(texto='OCR con error',alertas=['ocr_sin_cotejo_visual'],imagen_sha256='1'*64)
        check={'documento':'D1','ubicacion':'linea:1','sha256':docs[0]['sha256'],'imagen_sha256':'2'*64,
               'transcripcion':facts[0]['evidencias'][0]['cita'],'autor':'Simulación','tipo_autor':'asistente_ia','fecha':'2026-10-31T12:00:00-05:00'}
        with self.assertRaises(ValueError): aplicar_cotejos(docs,[check])
        check['imagen_sha256']='1'*64;aplicar_cotejos(docs,[check])
        self.assertEqual('OCR con error',b['texto_ocr_original'])
        self.assertEqual('cumple',next(r for r in evaluar(ficha(),docs,facts)['matriz'] if r['id']=='AC-03')['resultado'])

class OperacionV1(unittest.TestCase):
    def setUp(self):
        (ROOT/'data/pruebas-v1').mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=ROOT/'data/pruebas-v1');self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.case=self.root/'caso';iniciar(self.case,sintetico=True)
        guardar(self.case/'expediente.json',ficha());(self.case/'documentos/solicitud.txt').write_text('Solicitud sintética de ampliación de categoría. No es un documento real.',encoding='utf-8')
        self.memory=Memoria(self.root/'.oece')
    def test_entrega_nueva_repeticion_y_revision_preservada(self):
        first,reused=procesar(self.case,self.memory,False);self.assertFalse(reused)
        self.assertTrue((first/'informe.docx').exists());self.assertTrue((first/'matriz.csv').exists())
        self.assertEqual((first,True),procesar(self.case,self.memory,False))
        (self.case/'documentos/solicitud.txt').write_text('Cambio documentado posterior',encoding='utf-8')
        second,reused=procesar(self.case,self.memory,False)
        self.assertNotEqual(first,second);self.assertTrue((first/'entrega.json').exists())
    def test_entrega_alterada_no_se_reutiliza(self):
        folder,_=procesar(self.case,self.memory,False);(folder/'matriz.csv').write_text('alteración')
        with self.assertRaises(ValueError): procesar(self.case,self.memory,False)
    def test_documento_duplicado_conservado_sin_contar_doble(self):
        (self.case/'documentos/copia.txt').write_bytes((self.case/'documentos/solicitud.txt').read_bytes())
        result,_=preparar(self.case,self.memory,False)
        self.assertEqual(1,len(result['documentos']));self.assertEqual(2,len(result['instantanea']['archivos']))
    def test_estado_y_mensajes_sobreviven_a_reinicio(self):
        responder(self.case,self.memory,'estado')
        memory2=Memoria(self.memory.root)
        self.assertEqual(2,len(memory2.historial(sha(str(self.case.resolve()).encode()))['mensajes']))
    def test_vigilancia_espera_estabilidad_y_no_repite(self):
        (self.case/'LISTO').touch()
        self.assertEqual('esperando_archivos',observar_una_vez(self.root,self.memory,0,False)[0]['estado'])
        self.assertEqual('borrador',observar_una_vez(self.root,self.memory,0,False)[0]['estado'])
        self.assertEqual([],observar_una_vez(self.root,Memoria(self.memory.root),0,False))
    def test_archivo_parcial_no_se_procesa(self):
        (self.case/'documentos/carga.part').write_text('a')
        with self.assertRaises(ValueError): preparar(self.case,self.memory,False)
    def test_fallo_de_word_no_publica_entrega_parcial(self):
        with patch('oece.flujo.redactar',side_effect=RuntimeError('fallo simulado')):
            with self.assertRaises(RuntimeError): procesar(self.case,self.memory,False)
        self.assertEqual([],list((self.case/'salida').iterdir()))
    def test_dos_procesos_no_entran_al_mismo_expediente(self):
        with bloqueo(self.root/'lock'):
            with self.assertRaises(ValueError):
                with bloqueo(self.root/'lock'): pass
    def test_word_encabezado_y_tabla_leidos(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Texto principal');doc.sections[0].header.paragraphs[0].text='Identificador en encabezado'
        doc.add_table(rows=1,cols=1).cell(0,0).text='Contenido de tabla'
        path=self.case/'documentos/documento.docx';doc.save(path)
        result=leer_documento(path,self.root/'cache',False)
        text=' '.join(b['texto'] for b in result['fragmentos'])
        self.assertIn('Identificador en encabezado',text);self.assertIn('Contenido de tabla',text)
    def test_html_script_no_aparece_como_evidencia(self):
        p=self.case/'documentos/fuente.html';p.write_text('<p>Texto válido</p><script>orden maliciosa</script>',encoding='utf-8')
        result=leer_documento(p,self.root/'cache',False)
        text=' '.join(b['texto'] for b in result['fragmentos']);self.assertIn('Texto válido',text);self.assertNotIn('maliciosa',text)
    def test_ocr_timeout_queda_pendiente_y_puede_reanudar(self):
        from PIL import Image
        import subprocess
        p=self.case/'documentos/escaneo.png';Image.new('RGB',(100,100),'white').save(p)
        with patch('oece.lectura.ocr_imagen',side_effect=subprocess.TimeoutExpired('ocr',1)):
            first=leer_documento(p,self.root/'cache')
        self.assertTrue(first['reintentar']);self.assertIn('ocr_pendiente',first['fragmentos'][0]['alertas'])
        with patch('oece.lectura.ocr_imagen',return_value={'text':'Texto reconocido de prueba','image_sha256':'1'*64}):
            second=leer_documento(p,self.root/'cache')
        self.assertFalse(second['reintentar']);self.assertEqual('Texto reconocido de prueba',second['fragmentos'][0]['texto'])
    def test_xlsx_omite_celdas_solo_con_formato_y_no_valida_formulas(self):
        import zipfile
        p=self.case/'documentos/libro.xlsx'
        with zipfile.ZipFile(p,'w') as z:
            z.writestr('xl/worksheets/sheet1.xml','<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" s="1"/><c r="B1"><v>100</v></c><c r="C1"><f>B1+1</f><v>101</v></c></row></sheetData></worksheet>')
        result=leer_documento(p,self.root/'cache',False)
        self.assertEqual(2,len(result['fragmentos']))
        self.assertIn('formula_no_recalculada',result['fragmentos'][1]['alertas'])
    def test_norma_adicional_fuera_de_manifiesto_rechazada(self):
        from oece.biblioteca import validar_normas
        with self.assertRaises(ValueError):validar_normas([{'codigo':'N-1','archivo':'inventada.pdf','sha256':'0'*64,'url':'https://example.com'}])

class PanelV1(unittest.TestCase):
    def setUp(self):
        (ROOT/'data/pruebas-v1').mkdir(parents=True,exist_ok=True)
        self.temp=tempfile.TemporaryDirectory(dir=ROOT/'data/pruebas-v1');self.addCleanup(self.temp.cleanup)
        root=Path(self.temp.name);self.server=crear_servidor(root,Memoria(root/'.oece'),0)
        self.worker=threading.Thread(target=self.server.serve_forever,daemon=True);self.worker.start()
        self.base=f'http://127.0.0.1:{self.server.server_port}'
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.server.oece_pool.shutdown(wait=True);self.worker.join()
    def request(self,path,data=None,token=True,origin=None):
        headers={}
        if token: headers['X-OECE-Token']=self.server.oece_token
        if origin: headers['Origin']=origin
        request=Request(self.base+path,data=json.dumps(data).encode() if data is not None else None,headers=headers)
        with urlopen(request,timeout=10) as response:
            return json.loads(response.read())
    def test_creacion_subida_proceso_y_descarga(self):
        c=self.request('/api/nuevo',{'solicitante':'Ejemplo local'})['caso']
        self.request('/api/archivo',{'caso':c,'nombre':'doc.txt','contenido':base64.b64encode(b'Documento sintetico legible').decode()})
        self.request('/api/procesar',{'caso':c})
        for i in range(100):
            result=self.request('/api/casos')['casos'][0]
            if result['revisiones']: break
            time.sleep(.05)
        self.assertTrue(result['revisiones'])
        with urlopen(self.base+'/descargar?caso='+c+'&revision='+result['revisiones'][0]+'&archivo=informe.docx') as response:
            self.assertTrue(response.read().startswith(b'PK'))
    def test_sin_token_y_origen_externo_rechazados(self):
        with self.assertRaises(HTTPError) as e:self.request('/api/nuevo',{},token=False)
        self.assertEqual(403,e.exception.code)
        with self.assertRaises(HTTPError) as e:self.request('/api/nuevo',{},origin='https://ajeno.example')
        self.assertEqual(403,e.exception.code)
    def test_ruta_y_nombre_peligroso_rechazados(self):
        c=self.request('/api/nuevo',{})['caso']
        for name in ('../fuera.txt','C:\\fuera.txt','a.txt:flujo','CON.txt'):
            with self.assertRaises(HTTPError):self.request('/api/archivo',{'caso':c,'nombre':name,'contenido':'YQ=='})

if __name__=='__main__':unittest.main()
