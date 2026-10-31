"""Pruebas tecnicas sinteticas; no validan la matriz juridica del piloto."""
import hashlib
import importlib.util
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch
from contextlib import closing

from docx import Document

MODULE_PATH = Path(__file__).resolve().parents[1] / "scripts/indexar_normativa.py"
SPEC = importlib.util.spec_from_file_location("indexar_normativa", MODULE_PATH)
indexer = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(indexer)


class IndexTests(unittest.TestCase):
    def setUp(self):
        test_area = MODULE_PATH.parents[1] / "normativa/indices/pruebas"
        test_area.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=test_area)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        (self.root / "normativa").mkdir()
        self.source = self.root / "normativa/fuente.docx"
        document = Document()
        document.add_paragraph("Evaluación de inscripción y categoría")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Experiencia acreditada"
        table.cell(0, 1).text = "Consorcio y participación"
        nested = table.cell(0, 1).add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "Documento extranjero"
        document.save(self.source)
        self.item = {"name": self.source.name, "local_path": "normativa/fuente.docx",
                     "sha256": hashlib.sha256(self.source.read_bytes()).hexdigest(),
                     "bytes": self.source.stat().st_size,
                     "source_url": "https://www.gob.pe/fuente-sintetica-no-real",
                     "instrument_id": "sintetica"}
        self.catalog = {"normas": [{"id": "sintetica", "archivo": {"sha256": self.item["sha256"]},
                                   "vigencia": {"desde": None}, "validacion": {"nivel": "pendiente"},
                                   "procedimientos": ["ampliacion_consultores"]}]}
        (self.root / "normativa/catalogo.json").write_text(json.dumps(self.catalog), encoding="utf-8")
        self.output = self.root / "normativa/indices/biblioteca.sqlite"
        self.write_manifest([self.item])

    def write_manifest(self, items):
        (self.root / "manifest.json").write_text(json.dumps({"archivos": items}), encoding="utf-8")

    def build(self):
        return indexer.build_index(self.root, self.output, ("manifest.json",))

    def test_external_manifest_rejected_without_replacing_index(self):
        self.build()
        original = self.output.read_bytes()
        with self.assertRaisesRegex(ValueError, "Manifiesto fuera"):
            indexer.build_index(self.root, self.output, ("../fuera.json",))
        self.assertEqual(original, self.output.read_bytes())

    def test_search_preserves_nested_locator_hash_and_accent_insensitive_terms(self):
        self.build()
        results = indexer.search_index(self.output, "documento extranjero", "ampliacion_consultores")
        self.assertEqual(len(results), 1)
        self.assertIn("fila:1/celda:2", results[0]["ubicacion"]["ruta_bloque"])
        self.assertGreater(results[0]["ubicacion"]["ruta_bloque"].count("celda:"), 1)
        self.assertEqual(results[0]["documento"]["sha256"], self.item["sha256"])
        self.assertFalse(results[0]["apta_para_evaluacion"])
        self.assertTrue(indexer.search_index(self.output, "evaluacion inscripcion"))
        self.assertEqual(indexer.search_index(self.output, "evaluacion", "otro_procedimiento"), [])

    def test_modified_original_aborts_without_replacing_existing_index(self):
        self.build()
        previous = self.output.read_bytes()
        self.source.write_bytes(self.source.read_bytes() + b"modified")
        with self.assertRaisesRegex(ValueError, "Integridad"):
            self.build()
        self.assertEqual(previous, self.output.read_bytes())

    def test_manifest_cannot_read_outside_root(self):
        self.item["local_path"] = "../outside.docx"
        self.write_manifest([self.item])
        with self.assertRaisesRegex(ValueError, "fuera de la biblioteca"):
            self.build()

    def test_change_during_extraction_also_preserves_previous_index(self):
        self.build()
        previous = self.output.read_bytes()
        original_extract = indexer.extract

        def changing_extract(path):
            result = original_extract(path)
            path.write_bytes(path.read_bytes() + b"changed while reading")
            return result

        with patch.object(indexer, "extract", side_effect=changing_extract):
            with self.assertRaisesRegex(ValueError, "Integridad"):
                self.build()
        self.assertEqual(previous, self.output.read_bytes())

    def test_pdf_without_text_keeps_page_and_requires_ocr(self):
        from pypdf import PdfWriter
        source = self.root / "normativa/escaneado-sintetico.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        with source.open("wb") as stream:
            writer.write(stream)
        item = dict(self.item, name=source.name, local_path=source.relative_to(self.root).as_posix(),
                    bytes=source.stat().st_size, sha256=hashlib.sha256(source.read_bytes()).hexdigest())
        self.write_manifest([item])
        result = self.build()
        self.assertEqual(result["documentos"][0]["estado"], "requiere_ocr_o_revision")
        with closing(sqlite3.connect(self.output)) as connection:
            locator, flags = connection.execute("SELECT locator, flags FROM fragments").fetchone()
        self.assertEqual(json.loads(locator), {"pagina_pdf": 1})
        self.assertIn("requiere_ocr_o_revision", json.loads(flags))

    def test_same_content_keeps_provenance_without_duplicate_fragments(self):
        first = self.build()["fragmentos"]
        copy_path = self.root / "normativa/copia.docx"
        copy_path.write_bytes(self.source.read_bytes())
        copy_item = dict(self.item, name="copia.docx", local_path="normativa/copia.docx",
                         source_url="https://www.gob.pe/segunda-fuente-sintetica")
        self.write_manifest([self.item, copy_item])
        result = self.build()
        self.assertEqual(result["fragmentos"], first)
        match = indexer.search_index(self.output, "evaluacion")[0]
        self.assertEqual(len(match["documento"]["procedencias"]), 2)

    def test_unsupported_format_and_parser_error_are_reported(self):
        xls = self.root / "normativa/hoja.xls"
        xls.write_bytes(b"synthetic legacy format")
        bad = self.root / "normativa/corrupto.docx"
        bad.write_bytes(b"not an Office ZIP")
        items = []
        for path in (xls, bad):
            items.append(dict(self.item, name=path.name, local_path=path.relative_to(self.root).as_posix(),
                              bytes=path.stat().st_size, sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
        self.write_manifest(items)
        result = self.build()
        self.assertEqual([d["estado"] for d in result["documentos"]], ["formato_pendiente", "error_lectura"])
        self.assertEqual(result["fragmentos"], 0)

    def test_query_cannot_inject_sql_or_fts_operators(self):
        self.build()
        self.assertEqual(indexer.search_index(self.output, "\" OR 1=1; DROP TABLE fragments;--"), [])
        self.assertTrue(indexer.search_index(self.output, "experiencia"))
        with self.assertRaises(ValueError):
            indexer.search_index(self.output, " ")


if __name__ == "__main__":
    unittest.main()
