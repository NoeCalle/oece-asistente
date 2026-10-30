import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from crear_demo_expediente import create
from ejecutar_expediente import execute, load, prepare


class CaseFlowTests(unittest.TestCase):
    def setUp(self):
        area = ROOT / "data/pruebas-expedientes"
        area.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=area)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / "entrada"
        create(self.input)
        self.rules = ROOT / "rules/consultores/matriz.v0.1.json"

    def change(self, update):
        path = self.input / "expediente.json"
        data = load(path)
        update(data)
        path.write_text(json.dumps(data), encoding="utf-8")

    def test_end_to_end_and_idempotence(self):
        path, reused = execute(self.input, self.root / "salida")
        self.assertFalse(reused)
        result = load(path / "matriz.json")
        self.assertEqual(22, len(result["matriz"]))
        self.assertEqual({"no_determinable"}, {r["resultado"] for r in result["matriz"]})
        self.assertEqual(7, sum(len(r["evidencias"]) for r in result["matriz"]))
        from docx import Document
        doc = Document(path / "informe.docx")
        self.assertEqual(22, sum(len(t.rows) - 1 for t in doc.tables[1:]))
        self.assertEqual(3, len(load(path / "entrega.json")["estados"]))
        self.assertEqual((path, True), execute(self.input, self.root / "salida"))

    def test_modified_input_preserves_previous_revision(self):
        previous, _ = execute(self.input, self.root / "salida")
        with (self.input / "conformidad.txt").open("a", encoding="utf-8") as f:
            f.write("Nueva información sintética\n")
        current, reused = execute(self.input, self.root / "salida")
        self.assertNotEqual(previous, current)
        self.assertFalse(reused)
        self.assertTrue((previous / "entrega.json").is_file())
        self.assertNotEqual((previous / "originales/003.txt").read_bytes(),
                            (current / "originales/003.txt").read_bytes())

    def test_altered_output_is_rejected(self):
        path, _ = execute(self.input, self.root / "salida")
        (path / "matriz.json").write_text("alterado", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "alterada"):
            execute(self.input, self.root / "salida")

    def test_path_outside_case_rejected(self):
        self.change(lambda c: c["documentos"][0].update(archivo="../externo.txt"))
        with self.assertRaisesRegex(ValueError, "fuera"):
            prepare(self.input, self.rules)

    def test_real_case_rejected(self):
        self.change(lambda c: c.update(sintetico=False))
        with self.assertRaisesRegex(ValueError, "sintéticos"):
            prepare(self.input, self.rules)

    def test_unknown_control_rejected(self):
        self.change(lambda c: c["vinculos"].update({"AC-99": []}))
        with self.assertRaisesRegex(ValueError, "desconocido"):
            prepare(self.input, self.rules)

    def test_nonexistent_evidence_rejected(self):
        self.change(lambda c: c["vinculos"]["AC-01"][0].update(ubicacion="linea:999"))
        with self.assertRaisesRegex(ValueError, "inexistente"):
            prepare(self.input, self.rules)

    def test_duplicate_documents_identified(self):
        self.change(lambda c: c["documentos"].append({"id": "D4", "archivo": "contrato.txt"}))
        result = prepare(self.input, self.rules)
        self.assertEqual("D2", result["documentos"][-1]["duplicado_de"])

    def test_duplicate_id_rejected(self):
        self.change(lambda c: c["documentos"].append({"id": "D1", "archivo": "contrato.txt"}))
        with self.assertRaisesRegex(ValueError, "duplicado"):
            prepare(self.input, self.rules)

    def test_enabled_rules_rejected(self):
        data = load(self.rules)
        data["reglas"][0]["habilitada"] = True
        candidate = self.root / "reglas.json"
        candidate.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "deshabilitada"):
            prepare(self.input, candidate)

    def test_report_failure_leaves_no_partial_delivery(self):
        output = self.root / "salida"
        with patch("ejecutar_expediente.report", side_effect=RuntimeError("fallo simulado")):
            with self.assertRaises(RuntimeError):
                execute(self.input, output)
        self.assertEqual([], list(output.iterdir()))

    def test_pdf_and_word_locators(self):
        from docx import Document
        from pypdf import PdfWriter
        doc = Document()
        doc.add_paragraph("Documento ficticio Word")
        doc.save(self.input / "adjunto.docx")
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.write(self.input / "escaneo.pdf")
        self.change(lambda c: c["documentos"].extend([
            {"id": "D4", "archivo": "adjunto.docx"},
            {"id": "D5", "archivo": "escaneo.pdf"}]))
        result = prepare(self.input, self.rules)
        self.assertEqual("Documento ficticio Word", result["documentos"][3]["fragmentos"][0]["texto"])
        self.assertTrue(result["documentos"][4]["fragmentos"][0]["alertas"])


if __name__ == "__main__":
    unittest.main()
